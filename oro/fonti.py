"""Adattatori delle fonti dati. Ogni funzione dichiara SEMPRE la propria origine
e distingue «non disponibile» da «zero». Nessuna chiave API richiesta."""
from __future__ import annotations

import csv
import io
import json
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from datetime import date, datetime, timezone

ONCIA_TROY_G = 31.1034768  # grammi in un'oncia troy (definizione, non stima)
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) OroTrend/1.0"


@dataclass
class Lettura:
    """Esito di una lettura. `valore is None` significa NON RILEVATO, mai zero."""
    valore: float | None
    fonte: str
    momento: str | None = None      # data/ora dichiarata dalla fonte
    errore: str | None = None
    extra: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return self.valore is not None

    def __str__(self) -> str:
        if self.ok:
            return f"{self.valore} ({self.fonte}, {self.momento or 'senza data'})"
        return f"NON RILEVATO ({self.fonte}: {self.errore})"


# FRED strozza le raffiche di richieste ravvicinate: due serie chieste a distanza di
# millisecondi vanno in timeout, le stesse due a distanza di un secondo passano in 0,1s.
# Non e' un guasto ne' un blocco: e' cadenza. Si rispetta, non si aggira.
PAUSA_MINIMA_FRA_RICHIESTE = {"fred.stlouisfed.org": 1.5}
_ultima_richiesta: dict[str, float] = {}


def _attendi_il_turno(url: str) -> None:
    from urllib.parse import urlparse
    dominio = urlparse(url).netloc
    pausa = PAUSA_MINIMA_FRA_RICHIESTE.get(dominio)
    if not pausa:
        return
    trascorso = time.monotonic() - _ultima_richiesta.get(dominio, 0.0)
    if trascorso < pausa:
        time.sleep(pausa - trascorso)
    _ultima_richiesta[dominio] = time.monotonic()


def _get(url: str, tentativi: int = 3, timeout: int = 20) -> bytes:
    ultimo = None
    for n in range(tentativi):
        try:
            _attendi_il_turno(url)
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            ultimo = e
            if n < tentativi - 1:
                time.sleep(1.5 * (n + 1))
    raise ConnectionError(f"{type(ultimo).__name__}: {ultimo}")


# ------------------------------------------------------- apertura del mercato

def mercato_oro_aperto(momento=None) -> bool:
    """L'oro spot si scambia in continuo dalla domenica alle 22:00 UTC al venerdi' alle 22:00.

    ⛔ Serve perche' la fonte NON lo dice. A mercato chiuso gold-api.com continua a
    rispondere «updatedAtReadable: a few seconds ago» spostando in avanti il proprio
    orario mentre serve lo stesso identico prezzo dell'ultima chiusura. Fidandosi di
    quel campo si registrerebbero tre «misure» al sabato e tre alla domenica che sono
    tutte lo stesso numero di venerdi': non misure, copie.
    """
    m = (momento or datetime.now(timezone.utc)).astimezone(timezone.utc)
    giorno = m.weekday()          # 0 = lunedi' ... 4 = venerdi', 5 = sabato, 6 = domenica
    if giorno == 5:
        return False
    if giorno == 6:
        return m.hour >= 22
    if giorno == 4:
        return m.hour < 22
    return True


# ---------------------------------------------------------------- oro spot

def oro_spot_usd_oncia() -> Lettura:
    """Spot XAU/USD in dollari per oncia troy. Fonte primaria: gold-api.com."""
    fonte = "gold-api.com"
    try:
        d = json.loads(_get("https://api.gold-api.com/price/XAU"))
        p = float(d["price"])
        if not (500 < p < 50000):
            return Lettura(None, fonte, errore=f"valore fuori scala plausibile: {p}")
        return Lettura(p, fonte, momento=d.get("updatedAt"))
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def oro_futures_usd_oncia() -> Lettura:
    """Future COMEX GC=F. Fonte di RISCONTRO, non sostitutiva dello spot:
    il future quota con un premio/sconto rispetto allo spot."""
    fonte = "Yahoo Finance GC=F"
    try:
        d = json.loads(_get("https://query1.finance.yahoo.com/v8/finance/chart/GC=F?interval=1d&range=5d"))
        m = d["chart"]["result"][0]["meta"]
        return Lettura(float(m["regularMarketPrice"]), fonte,
                       momento=datetime.fromtimestamp(m["regularMarketTime"], timezone.utc).isoformat())
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def oro_storico_usd_oncia(intervallo: str = "5y") -> Lettura:
    """Serie storica giornaliera del future GC=F: {'AAAA-MM-GG': chiusura}."""
    fonte = f"Yahoo Finance GC=F ({intervallo})"
    try:
        d = json.loads(_get(
            f"https://query1.finance.yahoo.com/v8/finance/chart/GC=F?interval=1d&range={intervallo}"))
        r = d["chart"]["result"][0]
        chiusure = r["indicators"]["quote"][0]["close"]
        serie = {
            datetime.fromtimestamp(ts, timezone.utc).date().isoformat(): float(c)
            for ts, c in zip(r["timestamp"], chiusure) if c is not None
        }
        if not serie:
            return Lettura(None, fonte, errore="serie vuota")
        return Lettura(float(len(serie)), fonte, momento=max(serie), extra={"serie": serie})
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


# ------------------------------------------- oro mensile, Banca Mondiale

PAGINA_PINK_SHEET = "https://www.worldbank.org/en/research/commodity-markets"
# Ripiego se la pagina cambia forma. ⛔ Il codice nell'indirizzo cambia a ogni edizione:
# cablarlo e basta significherebbe servire dati fermi senza accorgersene.
PINK_SHEET_RIPIEGO = ("https://thedocs.worldbank.org/en/doc/"
                      "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
                      "CMO-Historical-Data-Monthly.xlsx")


def oro_mensile_banca_mondiale() -> Lettura:
    """Prezzo dell'oro in dollari per oncia troy, MENSILE, dal 1960.

    Fonte: World Bank Commodity Price Data («Pink Sheet»), licenza CC-BY 4.0 — uso
    commerciale e redistribuzione consentiti con attribuzione.

    ⛔ Perche' serve, ed e' la ragione per cui questa funzione esiste: la serie
    GIORNALIERA dell'oro non ha una fonte gratuita con termini utilizzabili. Yahoo la
    vieta esplicitamente («automated means… mobile application, data feed»), e FRED ha
    dovuto RIMUOVERE le serie LBMA per licenza (oggi rispondono 404). Questa e' l'unica
    storia lunga dell'oro che si possa ridistribuire.

    ⚠️ E' mensile, e per il legame con i fattori macro va benissimo: il segnale misurato
    vive proprio a orizzonte mensile (r = -0,60 contro -0,08 sul giornaliero), e con dati
    mensili le osservazioni sono davvero indipendenti — niente finestre sovrapposte.
    """
    fonte = "World Bank Pink Sheet (CC-BY 4.0)"
    try:
        import re
        try:
            pagina = _get(PAGINA_PINK_SHEET, timeout=40).decode("utf-8", errors="ignore")
            trovati = re.findall(r"https://[^\"']*CMO-Historical-Data-Monthly\.xlsx", pagina)
            url = trovati[0] if trovati else PINK_SHEET_RIPIEGO
        except Exception:
            url = PINK_SHEET_RIPIEGO
        grezzo = _get(url, timeout=90)

        import io as _io
        import openpyxl
        wb = openpyxl.load_workbook(_io.BytesIO(grezzo), read_only=True, data_only=True)
        sh = wb["Monthly Prices"]

        colonna = None
        for riga in sh.iter_rows(min_row=1, max_row=10, values_only=True):
            for j, c in enumerate(riga):
                if c and str(c).strip().lower() == "gold":
                    colonna = j
                    break
            if colonna is not None:
                break
        if colonna is None:
            return Lettura(None, fonte, errore="colonna «Gold» non trovata nel foglio mensile")

        serie: dict[str, float] = {}
        for riga in sh.iter_rows(min_col=1, max_col=colonna + 1, values_only=True):
            etichetta, valore = riga[0], riga[colonna]
            if not etichetta or not isinstance(valore, (int, float)):
                continue
            # Le righe utili hanno la forma «1960M01»: tutto il resto e' intestazione.
            testo = str(etichetta).strip()
            if len(testo) == 7 and testo[4] == "M" and testo[:4].isdigit() and testo[5:].isdigit():
                serie[f"{testo[:4]}-{testo[5:]}"] = round(float(valore), 4)

        if not serie:
            return Lettura(None, fonte, errore="nessuna osservazione mensile estratta")
        ultimo = max(serie)
        return Lettura(serie[ultimo], fonte, momento=ultimo,
                       extra={"serie": serie, "url": url, "unita": "USD per oncia troy"})
    except Exception as e:
        return Lettura(None, fonte, errore=f"{type(e).__name__}: {e}")


# ---------------------------------------------------------------- cambio

BCE_CAMBIO = "https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A"


def _leggi_csv_bce(testo: str) -> dict[str, float]:
    """La BCE pubblica DOLLARI PER EURO; a noi servono EURO PER DOLLARO.
    ⛔ Il verso e' l'errore piu' facile e il piu' silenzioso: 1,1652 e 0,8582 sono
    entrambi numeri plausibili, e sbagliarli falsa il prezzo del 35%."""
    fuori: dict[str, float] = {}
    for riga in csv.DictReader(io.StringIO(testo)):
        giorno, valore = riga.get("TIME_PERIOD"), riga.get("OBS_VALUE")
        if giorno and valore:
            try:
                usd_per_eur = float(valore)
                if usd_per_eur > 0:
                    fuori[giorno.strip()] = 1.0 / usd_per_eur
            except ValueError:
                continue
    return fuori


def eur_per_usd_bce() -> Lettura:
    """Cambio dalla BCE, fonte ufficiale senza intermediari."""
    fonte = "BCE (data-api.ecb.europa.eu)"
    try:
        testo = _get(f"{BCE_CAMBIO}?lastNObservations=1&format=csvdata", timeout=25).decode("utf-8")
        serie = _leggi_csv_bce(testo)
        if not serie:
            return Lettura(None, fonte, errore="nessuna osservazione nel CSV")
        g = max(serie)
        return Lettura(round(serie[g], 6), fonte, momento=g)
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def eur_per_usd_bce_storico(dal: str) -> Lettura:
    fonte = "BCE (data-api.ecb.europa.eu), serie"
    try:
        testo = _get(f"{BCE_CAMBIO}?startPeriod={dal}&format=csvdata", timeout=60).decode("utf-8")
        serie = {g: round(v, 6) for g, v in _leggi_csv_bce(testo).items()}
        if not serie:
            return Lettura(None, fonte, errore="serie vuota")
        return Lettura(float(len(serie)), fonte, momento=max(serie), extra={"serie": serie})
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def eur_per_usd() -> Lettura:
    """Quanti euro vale un dollaro.

    ⛔ Prima la BCE direttamente, poi frankfurter come ripiego. frankfurter e' un servizio
    di terzi che rispecchia gli stessi dati: il 09/09/2026 ha risposto 521 e 522 due volte
    di fila. Per un'app pubblica una dipendenza da un intermediario amatoriale, quando la
    fonte ufficiale e' interrogabile, e' un rischio che non ha ragione di esistere.
    """
    diretta = eur_per_usd_bce()
    if diretta.ok:
        return diretta
    fonte = f"frankfurter.dev (ripiego; BCE muta: {diretta.errore})"
    try:
        d = json.loads(_get("https://api.frankfurter.dev/v1/latest?base=USD&symbols=EUR"))
        v = float(d["rates"]["EUR"])
        if not (0.3 < v < 3.0):
            return Lettura(None, fonte, errore=f"valore fuori scala plausibile: {v}")
        return Lettura(v, fonte, momento=d.get("date"))
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def eur_per_usd_storico(dal: str) -> Lettura:
    """Serie storica del cambio USD->EUR da `dal` (AAAA-MM-GG) a oggi."""
    diretta = eur_per_usd_bce_storico(dal)
    if diretta.ok:
        return diretta
    fonte = f"frankfurter.dev (ripiego; BCE muta: {diretta.errore})"
    try:
        d = json.loads(_get(f"https://api.frankfurter.dev/v1/{dal}..?base=USD&symbols=EUR"))
        serie = {g: float(v["EUR"]) for g, v in d["rates"].items() if "EUR" in v}
        if not serie:
            return Lettura(None, fonte, errore="serie vuota")
        return Lettura(float(len(serie)), fonte, momento=max(serie), extra={"serie": serie})
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


# ---------------------------------------------------------------- conversione

def eur_grammo(usd_oncia: float, eur_su_usd: float, millesimi: int = 999) -> float:
    """Da dollari/oncia a euro/grammo. `millesimi`: 999 = oro fino 24k, 750 = 18k."""
    return usd_oncia * eur_su_usd / ONCIA_TROY_G * (millesimi / 1000.0)


# ---------------------------------------------------------------- macro FRED

SERIE_FRED = {
    "tasso_fed":        ("DFEDTARU", "Tasso obiettivo Fed, limite superiore", "%"),
    "tasso_reale_10a":  ("DFII10",   "Rendimento reale Treasury 10 anni (TIPS)", "%"),
    "inflazione_attesa":("T10YIE",   "Inflazione attesa a 10 anni (breakeven)", "%"),
    "nominale_10a":     ("DGS10",    "Rendimento nominale Treasury 10 anni", "%"),
    "dollaro":          ("DTWEXBGS", "Indice del dollaro, ponderato sul commercio", "indice"),
    "vix":              ("VIXCLS",   "VIX, volatilita' attesa S&P 500", "indice"),
    "inflazione_usa":   ("CPIAUCSL", "Indice prezzi al consumo USA", "indice"),
}


def fred(chiave: str) -> Lettura:
    """Scarica una serie FRED in CSV. Non richiede chiave API."""
    if chiave not in SERIE_FRED:
        raise KeyError(f"serie sconosciuta: {chiave}")
    codice, descrizione, unita = SERIE_FRED[chiave]
    fonte = f"FRED {codice}"
    try:
        testo = _get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={codice}").decode("utf-8")
        serie: dict[str, float] = {}
        for riga in csv.DictReader(io.StringIO(testo)):
            valori = list(riga.values())
            g, v = valori[0], valori[1]
            if v and v.strip() not in (".", ""):
                serie[g.strip()] = float(v)
        if not serie:
            return Lettura(None, fonte, errore="nessuna osservazione valida")
        ultimo = max(serie)
        return Lettura(serie[ultimo], fonte, momento=ultimo,
                       extra={"serie": serie, "descrizione": descrizione, "unita": unita, "codice": codice})
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


# ---------------------------------------------------------------- rischio geopolitico

def rischio_geopolitico() -> Lettura:
    """Indice GPR giornaliero di Caldara & Iacoviello (Federal Reserve Board).
    Misura la quota di articoli su guerre, tensioni e attentati in 10 grandi quotidiani.
    Fonte accademica, non commerciale: American Economic Review 2022, 112(4)."""
    fonte = "GPR daily (Caldara & Iacoviello, Federal Reserve Board)"
    url = "https://www.matteoiacoviello.com/gpr_files/data_gpr_daily_recent.xls"
    try:
        serie, media7, eventi = _leggi_gpr(_get(url, timeout=45))
        if not serie:
            return Lettura(None, fonte, errore="nessuna osservazione estratta dal file")
        ultimo = max(serie)
        return Lettura(serie[ultimo], fonte, momento=ultimo,
                       extra={"serie": serie, "media7": media7, "eventi_nominati": eventi})
    except Exception as e:
        return Lettura(None, fonte, errore=str(e))


def _leggi_gpr(grezzo: bytes) -> tuple[dict[str, float], dict[str, float], dict[str, str]]:
    """Legge il file .xls legacy (OLE2) del GPR con xlrd.
    Restituisce (indice giornaliero, media mobile 7 giorni, eventi nominati dagli autori)."""
    import xlrd  # dipendenza dichiarata: xlrd 2.x, legge SOLO .xls storico

    wb = xlrd.open_workbook(file_contents=grezzo)
    sh = wb.sheet_by_index(0)
    intest = [str(sh.cell_value(0, c)).strip().upper() for c in range(sh.ncols)]

    def colonna(*nomi: str) -> int | None:
        for n in nomi:
            if n in intest:
                return intest.index(n)
        return None

    i_giorno, i_gpr = colonna("DAY", "DATE"), colonna("GPRD", "GPR")
    i_ma7, i_evento = colonna("GPRD_MA7"), colonna("EVENT")
    if i_giorno is None or i_gpr is None:
        raise ValueError(f"colonne attese non trovate; presenti: {intest}")

    serie: dict[str, float] = {}
    media7: dict[str, float] = {}
    eventi: dict[str, str] = {}
    for r in range(1, sh.nrows):
        grezza = str(sh.cell_value(r, i_giorno)).strip()
        if grezza.endswith(".0"):
            grezza = grezza[:-2]
        if len(grezza) != 8 or not grezza.isdigit():
            continue
        g = f"{grezza[:4]}-{grezza[4:6]}-{grezza[6:]}"
        v = sh.cell_value(r, i_gpr)
        if isinstance(v, (int, float)) and v != "":
            serie[g] = float(v)
        if i_ma7 is not None:
            m = sh.cell_value(r, i_ma7)
            if isinstance(m, (int, float)) and m != "":
                media7[g] = float(m)
        if i_evento is not None:
            e = str(sh.cell_value(r, i_evento)).strip()
            if e and e not in ("0", "0.0"):
                eventi[g] = e
    return serie, media7, eventi
